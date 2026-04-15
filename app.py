import os
import json
import io
from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS
from dotenv import load_dotenv
from crewai import Agent, Task, Crew, Process
from openai import OpenAI

load_dotenv()

app = Flask(__name__)
CORS(app)

client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

try:
    from pypdf import PdfReader
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def extract_text(file_bytes, filename):
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "pdf" and HAS_PYPDF:
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if ext in ("docx", "doc") and HAS_DOCX:
        doc = DocxDocument(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)
    try:
        return file_bytes.decode("utf-8", errors="ignore")
    except Exception:
        return ""


FORMAT_GUIDE = {
    "instagram":       "Write an Instagram post: punchy opening, 3-4 sentences of body copy, clear CTA, 10-15 hashtags.",
    "facebook":        "Write a Facebook post: conversational, 100-200 words, question or CTA at the end. No hashtags.",
    "linkedin":        "Write a LinkedIn post: bold hook (never 'excited to share'), 3-5 short paragraphs, end with a question or CTA, 5 hashtags.",
    "twitter":         "Write 3 X (Twitter) variations each under 280 chars. Label: Tweet 1:, Tweet 2:, Tweet 3:",
    "google_ad":       "Write a Google Search Ad: 3 Headlines (max 30 chars), 2 Descriptions (max 90 chars), 2 Sitelink extensions. Label each.",
    "linkedin_ad":     "Write a LinkedIn Ad: Intro text (150 chars), Headline (70 chars), Description (100 chars), CTA. Also write a Message Ad version with subject + 300-word body.",
    "single_email":    "Write one marketing email: 2 subject line A/B options, preview text, 200-300 word body (hook/value/proof/CTA), PS line.",
    "email_campaign":   "Write a COMPLETE 3-email nurture sequence. Every email must be fully written — every sentence, every word, ready to send. DO NOT write outlines. Use this exact format:\n\nEMAIL 1 — AWARENESS\nSubject: [actual subject line]\nPreview Text: [actual preview text]\nBody:\n[Write full email — greeting, 4 complete paragraphs, sign-off. Min 200 words.]\nCTA: [CTA text]\n\n---\n\nEMAIL 2 — VALUE\nSubject: [actual subject line]\nPreview Text: [actual preview text]\nBody:\n[Write full email — greeting, 4 complete paragraphs, sign-off. Min 200 words.]\nCTA: [CTA text]\n\n---\n\nEMAIL 3 — CONVERSION\nSubject: [actual subject line]\nPreview Text: [actual preview text]\nBody:\n[Write full email — greeting, 4 complete paragraphs, sign-off. Min 200 words.]\nCTA: [CTA text]",    "blog":            "Write a full blog post: SEO title, meta description (155 chars), intro (100 words), 4-5 H2 sections (150-200 words each), conclusion with CTA. 900-1200 words total.",
    "blog":             "Write a COMPLETE fully written blog post — every word final, publish-ready. DO NOT write an outline. Format:\n\n## SEO TITLE\n[full title]\n\n## META DESCRIPTION\n[155-char description]\n\n## INTRODUCTION\n[3 full paragraphs]\n\n## [SECTION 1 HEADING]\n[4-5 full paragraphs]\n\n## [SECTION 2 HEADING]\n[4-5 full paragraphs]\n\n## [SECTION 3 HEADING]\n[4-5 full paragraphs]\n\n## [SECTION 4 HEADING]\n[4-5 full paragraphs]\n\n## [SECTION 5 HEADING]\n[3-4 full paragraphs]\n\n## CONCLUSION\n[2 full paragraphs + strong CTA]\n\nMinimum 1,200 words of fully written body content. No placeholders.",    "guide_pdf":        "Write a COMPLETE fully written lead-gen guide — every word final. DO NOT write a structure plan. Format:\n\n## GUIDE TITLE\n[full title]\n\n## SUBTITLE\n[subtitle]\n\n## EXECUTIVE SUMMARY\n[3 full paragraphs]\n\n## CHAPTER 1: [Title]\n[5-6 full paragraphs]\n\n## CHAPTER 2: [Title]\n[5-6 full paragraphs]\n\n## CHAPTER 3: [Title]\n[5-6 full paragraphs]\n\n## CHAPTER 4: [Title]\n[5-6 full paragraphs]\n\n## CHAPTER 5: [Title]\n[5-6 full paragraphs]\n\n## KEY TAKEAWAYS\n[6-8 bullet points as full sentences]\n\n## FINAL CTA\n[complete persuasive closing + call to action]\n\nMinimum 1,500 words. No placeholders.",
    "guide_pdf":        "Write a COMPLETE fully written lead-gen guide — every word final. DO NOT write a structure plan. Format:\n\n## GUIDE TITLE\n[full title]\n\n## SUBTITLE\n[subtitle]\n\n## EXECUTIVE SUMMARY\n[3 full paragraphs]\n\n## CHAPTER 1: [Title]\n[5-6 full paragraphs]\n\n## CHAPTER 2: [Title]\n[5-6 full paragraphs]\n\n## CHAPTER 3: [Title]\n[5-6 full paragraphs]\n\n## CHAPTER 4: [Title]\n[5-6 full paragraphs]\n\n## CHAPTER 5: [Title]\n[5-6 full paragraphs]\n\n## KEY TAKEAWAYS\n[6-8 bullet points as full sentences]\n\n## FINAL CTA\n[complete persuasive closing + call to action]\n\nMinimum 1,500 words. No placeholders.",}

def build_format_block(output_list):
    lines = []
    for fmt in output_list:
        if fmt == "images":
            continue
        guide = FORMAT_GUIDE.get(fmt, f"Write compelling content for {fmt}.")
        header = fmt.upper().replace("_", " ")
        lines.append(f"=== {header} ===\n{guide}")
    return "\n\n".join(lines)


def parse_sections(raw_text, output_list):
    sections = {}
    text_formats = [f for f in output_list if f != "images"]
    for fmt in text_formats:
        header = fmt.upper().replace("_", " ")
        marker = f"=== {header} ==="
        start = raw_text.find(marker)
        if start == -1:
            sections[fmt] = raw_text
            continue
        start += len(marker)
        next_markers = [
            raw_text.find(f"=== {o.upper().replace('_', ' ')} ===", start)
            for o in text_formats if o != fmt
        ]
        valid_nexts = [n for n in next_markers if n > start]
        end = min(valid_nexts) if valid_nexts else len(raw_text)
        sections[fmt] = raw_text[start:end].strip()
    return sections


def generate_images(ctx, research_snippet):
    # Generate detailed prompts via GPT-4o
    raw = client.chat.completions.create(
        model="gpt-4o",
        messages=[
            {"role": "system", "content": """You are a world-class commercial photography director.

Write 4 hyper-realistic image generation prompts for a professional marketing campaign.

Every prompt MUST:
- Describe a real photograph, not an illustration or digital art
- Include camera details: "shot on Sony A7R V" or "Canon EOS R5 mirrorless"
- Include lens details: "85mm f/1.4 lens", "50mm f/2.8", "24-70mm f/2.8"
- Include lighting: "golden hour sunlight", "softbox studio lighting", "natural window light", "overcast diffused light"
- Include depth of field: "shallow depth of field", "bokeh background", "sharp foreground"
- End every prompt with: "photorealistic, commercial photography, ultra-detailed, 8K resolution, no CGI, no illustration, real photograph"
- Be extremely specific about the subject, setting, colors, mood, and composition
- 3-5 sentences per prompt

Aspect ratios: "1:1" for social, "16:9" for hero/banner, "9:16" for stories/reels.

Return ONLY a JSON array: [{"label": "...", "aspect_ratio": "1:1", "prompt": "..."}]"""},
            {"role": "user", "content": f"Campaign brief:\n{ctx}\n\nResearch:\n{research_snippet}"},
        ],
        temperature=0.8,
    ).choices[0].message.content.strip().strip("```json").strip("```")

    try:
        prompts = json.loads(raw)
    except Exception:
        prompts = [
            {"label": "Social Square", "aspect_ratio": "1:1", "prompt": f"Professional marketing campaign photo. {ctx[:200]}. Shot on Sony A7R V, 85mm f/1.4, studio lighting, photorealistic, commercial photography, 4K."},
            {"label": "Hero Banner", "aspect_ratio": "16:9", "prompt": f"Wide cinematic hero banner for marketing. {ctx[:200]}. Golden hour lighting, shallow depth of field, Canon EOS R5, ultra-detailed, commercial quality."},
        ]

    import google.generativeai as genai
    genai.configure(api_key=os.getenv("GOOGLE_API_KEY"))

    results = []
    for p in prompts[:4]:
        try:
            model = genai.GenerativeModel("gemini-3.1-flash-image")
            response = model.generate_content(
                contents=[p["prompt"]],
                generation_config=genai.GenerationConfig(
                    response_mime_type="image/jpeg",
                )
            )
            # Extract base64 image and convert to data URL
            import base64
            img_data = response.candidates[0].content.parts[0].inline_data.data
            img_b64 = base64.b64encode(img_data).decode("utf-8")
            results.append({
                "label": p["label"],
                "prompt": p["prompt"],
                "url": f"data:image/jpeg;base64,{img_b64}",
            })
        except Exception as e:
            # Fallback to DALL-E 3 HD if Gemini fails
            try:
                size_map = {"1:1": "1024x1024", "16:9": "1792x1024", "9:16": "1024x1792"}
                size = size_map.get(p.get("aspect_ratio", "1:1"), "1024x1024")
                img = client.images.generate(
                    model="dall-e-3",
                    prompt=p["prompt"],
                    size=size,
                    quality="hd",
                    style="natural",
                    n=1,
                )
                results.append({"label": p["label"], "prompt": p["prompt"], "url": img.data[0].url})
            except Exception as e2:
                results.append({"label": p["label"], "prompt": p["prompt"], "url": "", "error": str(e2)})
    return results


@app.route("/generate", methods=["POST"])
def generate():
    topic         = request.form.get("topic", "").strip()
    business_name = request.form.get("business_name", "").strip()
    business_type = request.form.get("business_type", "").strip()
    audience      = request.form.get("audience", "").strip()
    tone          = request.form.get("tone", "").strip()
    key_message   = request.form.get("key_message", "").strip()
    outputs_raw   = request.form.get("outputs", "[]")

    try:
        output_list = json.loads(outputs_raw)
    except Exception:
        return jsonify({"success": False, "error": "Invalid outputs format"}), 400

    if not output_list:
        return jsonify({"success": False, "error": "Select at least one output format"}), 400

    doc_text = ""
    for f in request.files.getlist("documents"):
        raw = f.read()
        doc_text += f"\n--- {f.filename} ---\n" + extract_text(raw, f.filename)

    ctx = f"""BUSINESS: {business_name} ({business_type})
CAMPAIGN TOPIC: {topic}
TARGET AUDIENCE: {audience}
TONE / VOICE: {tone}
KEY MESSAGE: {key_message or 'Not specified'}"""
    if doc_text:
        ctx += f"\n\nREFERENCE DOCUMENTS:\n{doc_text[:4000]}"

    formats_str  = ", ".join([f for f in output_list if f != "images"])
    format_block = build_format_block(output_list)

    researcher = Agent(
        role="Market Research Specialist",
        goal="Analyze business briefs and uncover audience insights, content angles, and messaging hooks",
        backstory="Senior market researcher with 15 years of B2B and B2C experience. Expert at identifying what makes target audiences take action.",
        verbose=True,
        allow_delegation=False,
    )
    strategist = Agent(
        role="Content Strategist",
        goal="Build a concise channel-specific content strategy based on research findings",
        backstory="Content strategist who has launched 500+ campaigns. Expert at matching message to channel.",
        verbose=True,
        allow_delegation=False,
    )
    copywriter = Agent(
        role="Expert Multi-Channel Copywriter",
        goal="Write compelling, publish-ready copy for every requested marketing format",
        backstory="Award-winning copywriter specializing in digital marketing across social, email, ads, and long-form.",
        verbose=True,
        allow_delegation=False,
    )
    editor = Agent(
        role="Senior Content Editor",
        goal="Refine and polish all copy to match the brand voice and ensure quality",
        backstory="Meticulous editor with a sharp eye for tone, clarity, and brand consistency.",
        verbose=True,
        allow_delegation=False,
    )

    research_task = Task(
        description=f"Analyze this brief and produce a research summary: key audience pain points, 3-5 content angles, competitive notes, and top messaging hooks.\n\nBRIEF:\n{ctx}",
        agent=researcher,
    )
    strategy_task = Task(
        description=f"Using the research, create a content strategy for these formats: {formats_str}. Include one core message, CTA per channel, and channel notes.\n\nBRIEF:\n{ctx}",
        agent=strategist,
    )
    copy_task = Task(
        description=f"""Write publish-ready content for every format below.
Start EACH format with its exact header (e.g. === INSTAGRAM ===). Do NOT skip any.

BRIEF:
{ctx}

FORMATS TO WRITE:
{format_block}
""",
        agent=copywriter,
    )
    edit_task = Task(
        description=f"Review and refine all the copy. Match the '{tone}' tone. Keep all === FORMAT === headers exactly intact. Return the full polished content.",
        agent=editor,
    )

    crew = Crew(
        agents=[researcher, strategist, copywriter, editor],
        tasks=[research_task, strategy_task, copy_task, edit_task],
        process=Process.sequential,
        verbose=True,
    )

    try:
        result = crew.kickoff()
        raw_output = str(result)
    except Exception as e:
        return jsonify({"success": False, "error": "Pipeline failed", "detail": str(e)}), 500

    sections = parse_sections(raw_output, output_list)

    if "images" in output_list:
        sections["images"] = generate_images(ctx, raw_output[:800])

    return jsonify({
        "success": True,
        "inputs": {
            "topic": topic, "business_name": business_name,
            "business_type": business_type, "audience": audience,
            "tone": tone, "outputs": output_list,
        },
        "sections": sections,
    })


@app.route("/")
def index():
    return send_from_directory(".", "index.html")


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8080, debug=True)